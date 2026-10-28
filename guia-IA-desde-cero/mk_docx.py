import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches

SRC, OUT = sys.argv[1], sys.argv[2]
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(6)

def add_runs(par, text, base_bold=False):
    for p in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not p: continue
        if p.startswith("**") and p.endswith("**") and len(p) > 4:
            r = par.add_run(p[2:-2]); r.bold = True
        elif p.startswith("`") and p.endswith("`") and len(p) > 2:
            r = par.add_run(p[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10)
        else:
            r = par.add_run(p)
        if base_bold: r.bold = True

lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
while i < len(lines):
    s = lines[i].strip()
    if s.startswith("```"):
        i += 1; buf = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            buf.append(lines[i]); i += 1
        i += 1
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
        r = p.add_run("\n".join(buf)); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        continue
    if s.startswith("|") and i + 1 < len(lines) and set(lines[i+1].strip()) <= set("|-: "):
        header = [c.strip() for c in s.strip("|").split("|")]; i += 2; rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i += 1
        t = doc.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]; c.text = ""; add_runs(c.paragraphs[0], h, base_bold=True)
        for row in rows:
            cells = t.add_row().cells
            for j in range(len(header)):
                cells[j].text = ""; add_runs(cells[j].paragraphs[0], row[j] if j < len(row) else "")
        doc.add_paragraph(); continue
    if s.startswith("# "): doc.add_heading(s[2:], level=0)
    elif s.startswith("### "): doc.add_heading(s[4:], level=3)
    elif s.startswith("## "): doc.add_heading(s[3:], level=1)
    elif s.startswith("> "):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3); add_runs(p, s[2:])
        for r in p.runs: r.italic = True
    elif s.startswith("- [ ] "): doc.add_paragraph(s[6:], style="List Bullet")
    elif s.startswith("- "): doc.add_paragraph(s[2:], style="List Bullet")
    elif re.match(r"^\d+\.\s", s): doc.add_paragraph(re.sub(r"^\d+\.\s", "", s), style="List Number")
    elif s == "---": doc.add_paragraph()
    elif s: p = doc.add_paragraph(); add_runs(p, s)
    i += 1
doc.save(OUT)
print("OK ->", OUT)
